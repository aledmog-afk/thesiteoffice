#!/usr/bin/env python3
"""
Extracts structured content from the 30 supplied Toolbox Talk .docx files
into a JSON manifest. Read-only against the source documents: never
rewrites or "improves" wording, and flags (rather than guesses at) any
document whose structure doesn't match the expected shape.

Expected structure (verified identical across all 30 source documents
before writing this parser):
  Table 0: title table (row0/col0 = "TOOLBOX TALK\n<Title>")
  Table 1: meta table (Duration / Audience / Conducted By / Date & Location)
    -- NOT imported as content: Conducted By / Date & Location are
    -- delivery-time placeholders ([NAME & POSITION], [DATE], [SITE NAME])
    -- that this app's own toolbox_talks.delivered_by/started_at/project_id
    -- already provide natively.
  Heading "INTRODUCTION — READ THIS OUT" -> next paragraph(s) = introduction
  Table 2: KEY MESSAGE (row0/col0 = "KEY MESSAGE\n<text>")
  Heading "KEY POINTS TO COVER" -> following List Paragraph bullets
  Heading "WHAT TO LOOK FOR ON SITE" -> next table, 2 cols
    (col0 = "GOOD PRACTICE" + items, col1 = "WARNING SIGNS" + items)
  Heading "DISCUSSION QUESTIONS — ASK THE TEAM" -> following bullets
  Heading "KEY TAKEAWAYS — WHAT WE ALL AGREE TO DO" -> next table, 1x1,
    each paragraph in the cell = one takeaway
  Heading "PRE-TASK SAFETY CHECK" -> next table, first column (minus its
    own "Check Item" header row) = the checklist items
  Heading "ATTENDANCE REGISTER" -> NOT imported as content at all: this is
    the paper sign-in sheet this app's own digital attendee/signature
    system replaces.
  Trailing paragraph "Document produced by The Site Office..." -> generic
    boilerplate, identical across all 30 documents (verified) -- not
    document-specific, so NOT used as source_reference (no per-document
    version/revision is stated anywhere in these documents).
"""
import docx
import glob
import json
import os
import re
import sys

EXPECTED_HEADINGS = [
    "INTRODUCTION — READ THIS OUT",
    "KEY POINTS TO COVER",
    "WHAT TO LOOK FOR ON SITE",
    "DISCUSSION QUESTIONS — ASK THE TEAM",
    "KEY TAKEAWAYS — WHAT WE ALL AGREE TO DO",
    "PRE-TASK SAFETY CHECK",
    "ATTENDANCE REGISTER",
]


def para_texts_between(paragraphs, start_idx, headings_set):
    """Non-empty paragraph texts after start_idx, up to (not including)
    the next heading paragraph or end of document."""
    out = []
    for p in paragraphs[start_idx + 1:]:
        t = p.text.strip()
        if not t:
            continue
        if t in headings_set:
            break
        out.append(t)
    return out


def extract_one(path):
    d = docx.Document(path)
    paragraphs = d.paragraphs
    heading_idx = {}
    for i, p in enumerate(paragraphs):
        t = p.text.strip()
        if t in EXPECTED_HEADINGS:
            heading_idx.setdefault(t, i)

    issues = []
    missing = [h for h in EXPECTED_HEADINGS if h not in heading_idx]
    if missing:
        issues.append(f"missing heading(s): {missing}")

    if len(d.tables) < 6:
        issues.append(f"expected at least 6 tables, found {len(d.tables)}")
        return None, issues

    # Table 0: title
    title_cell_text = d.tables[0].rows[0].cells[0].text.strip()
    title_lines = [l.strip() for l in title_cell_text.split("\n") if l.strip()]
    if len(title_lines) < 2 or title_lines[0].upper() != "TOOLBOX TALK":
        issues.append(f"unexpected title table shape: {title_lines!r}")
        title = title_lines[-1] if title_lines else None
    else:
        title = title_lines[1]

    # Table 1: duration / audience
    duration_label = None
    audience = None
    try:
        meta_row = d.tables[1].rows[0]
        cells = [c.text.strip() for c in meta_row.cells]
        # ['Duration', '10–15 minutes', 'Audience', 'All site operatives']
        if len(cells) >= 4 and cells[0].lower() == "duration":
            duration_label = cells[1]
        if len(cells) >= 4 and cells[2].lower() == "audience":
            audience = cells[3]
    except Exception as e:
        issues.append(f"could not read meta table: {e}")

    # Introduction
    introduction = None
    if "INTRODUCTION — READ THIS OUT" in heading_idx:
        parts = para_texts_between(paragraphs, heading_idx["INTRODUCTION — READ THIS OUT"], set(EXPECTED_HEADINGS))
        introduction = "\n\n".join(parts) if parts else None
        if not introduction:
            issues.append("empty introduction section")

    # Key message (table 2)
    key_message = None
    try:
        km_text = d.tables[2].rows[0].cells[0].text.strip()
        km_lines = [l.strip() for l in km_text.split("\n") if l.strip()]
        if km_lines and km_lines[0].upper() == "KEY MESSAGE":
            key_message = "\n\n".join(km_lines[1:])
        else:
            key_message = km_text
            issues.append("KEY MESSAGE table label not found as expected")
    except Exception as e:
        issues.append(f"could not read key message table: {e}")

    # Key points
    key_points = []
    if "KEY POINTS TO COVER" in heading_idx:
        key_points = para_texts_between(paragraphs, heading_idx["KEY POINTS TO COVER"], set(EXPECTED_HEADINGS))
        if not key_points:
            issues.append("no key points found")

    # Good practice / warning signs (table 3, 2 columns)
    good_practice, warning_signs = [], []
    try:
        gp_table = d.tables[3]
        col0 = [p.text.strip() for p in gp_table.rows[0].cells[0].paragraphs if p.text.strip()]
        col1 = [p.text.strip() for p in gp_table.rows[0].cells[1].paragraphs if p.text.strip()]
        if col0 and col0[0].upper() == "GOOD PRACTICE":
            good_practice = col0[1:]
        else:
            issues.append("GOOD PRACTICE column header not found as expected")
        if col1 and col1[0].upper() == "WARNING SIGNS":
            warning_signs = col1[1:]
        else:
            issues.append("WARNING SIGNS column header not found as expected")
    except Exception as e:
        issues.append(f"could not read good practice/warning signs table: {e}")

    # Discussion questions
    discussion_questions = []
    if "DISCUSSION QUESTIONS — ASK THE TEAM" in heading_idx:
        discussion_questions = para_texts_between(paragraphs, heading_idx["DISCUSSION QUESTIONS — ASK THE TEAM"], set(EXPECTED_HEADINGS))
        if not discussion_questions:
            issues.append("no discussion questions found")

    # Key takeaways (table 4, 1x1, one paragraph per takeaway)
    key_takeaways = []
    try:
        kt_cell = d.tables[4].rows[0].cells[0]
        key_takeaways = [p.text.strip() for p in kt_cell.paragraphs if p.text.strip()]
        if not key_takeaways:
            issues.append("no key takeaways found")
    except Exception as e:
        issues.append(f"could not read key takeaways table: {e}")

    # Pre-task checks (table 5, first column, minus header row)
    pre_task_checks = []
    try:
        check_table = d.tables[5]
        header = [c.text.strip() for c in check_table.rows[0].cells]
        if not header or header[0].lower() != "check item":
            issues.append("pre-task check table header not found as expected")
        for row in check_table.rows[1:]:
            item = row.cells[0].text.strip()
            if item:
                pre_task_checks.append(item)
        if not pre_task_checks:
            issues.append("no pre-task checks found")
    except Exception as e:
        issues.append(f"could not read pre-task check table: {e}")

    # Images
    image_count = sum(1 for r in d.part.rels.values() if "image" in r.reltype)

    record = {
        "source_filename": os.path.basename(path),
        "title": title,
        "duration_label": duration_label,
        "audience": audience,
        "introduction": introduction,
        "key_message": key_message,
        "key_points": key_points,
        "site_observations": {
            "good_practice": good_practice,
            "warning_signs": warning_signs,
        },
        "discussion_questions": discussion_questions,
        "key_takeaways": key_takeaways,
        "pre_task_checks": pre_task_checks,
        "image_count": image_count,
    }
    return record, issues


def main():
    files = sorted(glob.glob("docs/*.docx"))
    results = []
    all_issues = {}
    for f in files:
        record, issues = extract_one(f)
        if issues:
            all_issues[os.path.basename(f)] = issues
        if record is not None:
            results.append(record)

    # Stable template_key derived from filename, e.g. TBT-01-WorkingAtHeight.docx -> TBT-01
    for r in results:
        m = re.match(r"(TBT-\d+)", r["source_filename"])
        r["template_key"] = m.group(1) if m else r["source_filename"]

    results.sort(key=lambda r: r["template_key"])

    with open("manifest.json", "w") as f:
        json.dump(results, f, indent=2, ensure_ascii=False)

    print(f"Documents found:        {len(files)}")
    print(f"Successfully extracted: {len(results)}")
    print(f"Documents with issues:  {len(all_issues)}")
    if all_issues:
        print("\n--- ISSUES (flagged for manual review, NOT silently skipped) ---")
        for fname, issues in all_issues.items():
            print(f"  {fname}:")
            for i in issues:
                print(f"    - {i}")

    # Sanity: every record has non-empty content in every section
    empty_field_report = []
    for r in results:
        for field in ["title", "introduction", "key_message"]:
            if not r.get(field):
                empty_field_report.append(f"{r['source_filename']}: empty {field}")
        for field in ["key_points", "discussion_questions", "key_takeaways", "pre_task_checks"]:
            if not r.get(field):
                empty_field_report.append(f"{r['source_filename']}: empty {field}")
        so = r.get("site_observations", {})
        if not so.get("good_practice") or not so.get("warning_signs"):
            empty_field_report.append(f"{r['source_filename']}: empty site_observations")
    if empty_field_report:
        print("\n--- EMPTY FIELD WARNINGS ---")
        for w in empty_field_report:
            print(" ", w)
    else:
        print("\nNo empty-field warnings — every record has content in every expected section.")

    return 0 if not all_issues else 1


if __name__ == "__main__":
    sys.exit(main())
